import pandas as pd
from openpyxl import load_workbook
from src.infrastructure.schedule_exporter import ScheduleExporter
from src.scheduling.time_model import TimeModel


def test_excel_and_csv_preserve_detail_contract(tmp_path):
    assignments = {'BIO-G1': ('A1', 1, 1260, 1320)}
    exporter = ScheduleExporter(TimeModel.default())
    csv_path, excel_path = tmp_path / 'schedule.csv', tmp_path / 'schedule.xlsx'
    exporter.to_csv(assignments, str(csv_path), course_name_by_code={'BIO': 'Biología'})
    exporter.to_excel(assignments, str(excel_path), course_name_by_code={'BIO': 'Biología'})
    csv = pd.read_csv(csv_path)
    excel = pd.read_excel(excel_path, sheet_name='Asignaciones')
    pd.testing.assert_frame_equal(csv, excel)
    assert list(csv.columns) == ['Código Curso', 'Nombre Curso', 'Grupo', 'Aula', 'Día', 'Hora Inicio', 'Hora Fin']
    assert csv.iloc[0].to_dict() == {
        'Código Curso': 'BIO', 'Nombre Curso': 'Biología', 'Grupo': 'BIO-G1',
        'Aula': 'A1', 'Día': 'Lunes', 'Hora Inicio': '21:00', 'Hora Fin': '22:00'}
    workbook = load_workbook(excel_path)
    assert len(workbook.sheetnames) == 3
    assert any('21:30' == cell.value for sheet in workbook for row in sheet for cell in row)
